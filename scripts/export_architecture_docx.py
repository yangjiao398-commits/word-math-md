"""Write the Word-to-App data-flow document."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "从Word上传到App刷题.docx"
FONT = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_B = Path(r"C:\Windows\Fonts\msyhbd.ttc")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_B if bold and FONT_B.is_file() else FONT
    return ImageFont.truetype(str(path), size)


def box(draw: ImageDraw.ImageDraw, xy, text: str, fill, outline="#1f4e79") -> None:
    x0, y0, x1, y1 = xy
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=2)
    lines = text.split("\n")
    f = font(22, bold=True)
    heights = [draw.textbbox((0, 0), line, font=f)[3] for line in lines]
    total = sum(heights) + 6 * (len(lines) - 1)
    y = y0 + (y1 - y0 - total) / 2
    for line, h in zip(lines, heights):
        w = draw.textbbox((0, 0), line, font=f)[2]
        draw.text(((x0 + x1 - w) / 2, y), line, fill="#102033", font=f)
        y += h + 6


def arrow(draw: ImageDraw.ImageDraw, x: int, y0: int, y1: int) -> None:
    draw.line((x, y0, x, y1 - 10), fill="#3d6ea8", width=4)
    draw.polygon([(x - 8, y1 - 12), (x + 8, y1 - 12), (x, y1)], fill="#3d6ea8")


def draw_flow(path: Path) -> None:
    w, h = 1600, 1180
    img = Image.new("RGB", (w, h), "#f7fafc")
    d = ImageDraw.Draw(img)
    title = font(36, bold=True)
    d.text((40, 28), "从 Word 上传到 App 刷题", fill="#102033", font=title)
    sub = font(20)
    d.text((40, 78), "管理端写入 Supabase，App 只读 /api/app/*", fill="#4a6072", font=sub)

    box(d, (480, 130, 1120, 230), "管理端  http://127.0.0.1:3010\n上传「解析版」.docx 或 .md", "#d6e8f8")
    arrow(d, 800, 230, 280)
    box(d, (80, 280, 360, 400), "1. OLE 公式 → LaTeX\nole_to_latex.py", "#e8f2ea")
    box(d, (400, 280, 680, 400), "2. Word 预处理\npreprocess_docx", "#e8f2ea")
    box(d, (720, 280, 1040, 400), "3. Node/TS 转 Markdown\ngaokao_docx + KaTeX", "#e8f2ea")
    box(d, (1080, 280, 1520, 400), "4. 拆题入库\nexam_bank.py", "#e8f2ea")
    for x in (360, 680, 1040):
        d.line((x, 340, x + 40, 340), fill="#3d6ea8", width=4)
        d.polygon([(x + 28, 332), (x + 28, 348), (x + 42, 340)], fill="#3d6ea8")

    arrow(d, 800, 400, 470)
    box(d, (180, 470, 760, 600), "Supabase PostgreSQL\npapers / questions / knowledge\napp_users / memberships", "#fff4d6")
    box(d, (840, 470, 1420, 600), "Supabase Storage\nexam-assets\n试卷图片、知识点 PPT", "#fff4d6")

    arrow(d, 470, 600, 680)
    arrow(d, 1130, 600, 680)
    box(d, (80, 680, 1520, 820), "App（math-agent-app）  登录 JWT → 校验会员 → 拉试卷 / 专项训练\nGET /api/app/papers  →  GET .../preview（KaTeX HTML）→  WebView 刷题", "#fde8e4")
    d.text(
        (80, 860),
        "可选：拍照阅卷  POST /api/app/papers/{id}/grade-sheet  →  OCR  →  answer_sheets",
        fill="#102033",
        font=font(22),
    )
    d.text(
        (80, 910),
        "可选：知识点 PPT  管理端上传 courseware  →  题目/知识点页浏览",
        fill="#102033",
        font=font(22),
    )
    d.text(
        (80, 980),
        "设计要点：Word 入库只走管理端 /api/exam-bank/*；App 不调用管理接口。",
        fill="#1f4e79",
        font=font(22, bold=True),
    )
    img.save(path)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 14, bold=True, color=(31, 78, 121))


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, 11)


def shade_header(row):
    for cell in row.cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): "1F4E79", qn("w:val"): "clear"})
        tcPr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    shade_header(table.rows[0])
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.rows[r].cells[c].text = value
            for p in table.rows[r].cells[c].paragraphs:
                for run in p.runs:
                    set_run_font(run, 9)
    doc.add_paragraph()


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    png = OUT.with_suffix(".flow.png")
    draw_flow(png)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("word-math-md\n从 Word 上传到 App 刷题")
    set_run_font(run, 20, bold=True, color=(31, 78, 121))
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("数据流与 /api/* 路由归属（管理端 vs App）")
    set_run_font(r2, 12, color=(74, 96, 114))

    add_heading(doc, "一、完整数据流", 1)
    add_para(doc, "管理端把解析版 Word 转成题目并写入 Supabase。App 只通过 /api/app/* 读取已入库内容，并用 WebView 渲染 KaTeX。")
    doc.add_picture(str(png), width=Cm(17.5))

    add_heading(doc, "分步说明", 2)
    add_table(
        doc,
        ["阶段", "谁触发", "发生了什么", "数据落在哪"],
        [
            ["1. 上传", "管理员在 /", "选择带「解析版」的 .docx，或直传 .md", "临时目录 word-math-md-uploads"],
            ["2. 转换", "batch_import", "OLE → 预处理 → Node 转 MD → 解析题目", "本地 .md 与图片"],
            ["3. 入库", "exam_bank", "写入试卷、题目、选项；图片进 Storage", "PostgreSQL + exam-assets"],
            ["4. 登录", "App 用户", "短信验证码 → JWT → 查会员", "app_users / app_memberships"],
            ["5. 看卷", "App 试卷页", "列表 → 预览 HTML → WebView", "只读数据库 + CDN 图片"],
            ["6. 专项", "App 练习页", "按知识点选题并组卷预览", "只读数据库"],
            ["7. 阅卷", "App 拍照", "OCR 后写入答卷", "answer_sheets"],
            ["8. 课件", "知识点维护", "上传 PPT，题目页按知识点浏览", "knowledge_point_courseware + Storage"],
        ],
    )

    add_heading(doc, "二、路由归属约定", 1)
    add_para(doc, "管理端：页面 /、/knowledge、/practice 使用，无 JWT，依赖本机信任。")
    add_para(doc, "App：math-agent-app 专用。内容接口需要 JWT，并且需要有效会员（否则 HTTP 402）。")
    add_para(doc, "回调/公共：支付平台或无需登录的接口。")

    add_heading(doc, "2.1 转换工具（管理端）", 2)
    add_table(
        doc,
        ["方法", "路径", "用途"],
        [
            ["POST", "/api/open-local", "打开本地文件"],
            ["POST", "/api/preprocess", "Word 预处理"],
            ["GET", "/api/preprocessed/{job}/{name}", "下载预处理结果"],
            ["POST", "/api/ole-to-latex", "OLE 公式转 LaTeX"],
            ["GET", "/api/ole-to-latex/{job}/{name}", "下载 OLE 结果"],
            ["POST", "/api/inspect", "检查 docx 结构"],
            ["POST", "/api/convert", "单文件转换预览"],
            ["GET", "/api/download/{name}", "下载 ZIP"],
            ["POST", "/api/parse-markdown", "解析 Markdown 预览"],
        ],
    )

    add_heading(doc, "2.2 题库管理（管理端 /api/exam-bank/*）", 2)
    add_table(
        doc,
        ["方法", "路径", "用途"],
        [
            ["GET", "/api/exam-bank/status", "是否已配置 Supabase"],
            ["POST", "/api/exam-bank/import-markdown", "导入 .md 入库"],
            ["POST", "/api/exam-bank/import-docx", "Word 上传入库（主路径）"],
            ["GET", "/api/exam-bank/papers", "试卷列表"],
            ["PATCH", "/api/exam-bank/papers/{paper_id}", "改学期、省份等元数据"],
            ["DELETE", "/api/exam-bank/papers/{paper_id}", "删除试卷及关联数据"],
            ["GET", "/api/exam-bank/papers/{paper_id}", "试卷详情"],
            ["GET", "/api/exam-bank/papers/{paper_id}/preview", "预览（含答案/解析）"],
            ["POST", "/api/exam-bank/questions-by-knowledge", "按知识点查题"],
            ["POST", "/api/exam-bank/practice/preview", "专项训练预览"],
            ["GET", "/api/exam-bank/knowledge-points", "知识点列表"],
            ["POST", "/api/exam-bank/knowledge-points", "新增知识点"],
            ["POST", "/api/exam-bank/knowledge-points/bulk", "批量导入"],
            ["POST", "/api/exam-bank/knowledge-points/seed-gaokao", "灌入高考知识点"],
            ["PUT", "/api/exam-bank/knowledge-points/{code}", "更新知识点"],
            ["DELETE", "/api/exam-bank/knowledge-points/{code}", "删除知识点"],
            ["PUT", "/api/exam-bank/questions/{question_id}/knowledge-points", "绑定题目知识点"],
            ["GET", "/api/exam-bank/knowledge-courseware", "按多个知识点查 PPT"],
            ["GET", "/api/exam-bank/knowledge-points/{code}/courseware", "某知识点的 PPT 列表"],
            ["POST", "/api/exam-bank/knowledge-points/{code}/courseware", "上传 PPT/PPTX"],
            ["DELETE", "/api/exam-bank/knowledge-courseware/{item_id}", "删除课件"],
            ["GET", "/api/exam-bank/ocr-status", "OCR 引擎状态"],
            ["GET", "/api/exam-bank/papers/{paper_id}/answer-sheets", "答卷记录"],
            ["GET", "/api/exam-bank/answer-sheets/{sheet_id}", "单份答卷"],
            ["POST", "/api/exam-bank/papers/{paper_id}/grade-sheet", "上传照片阅卷"],
            ["POST", "/api/exam-bank/papers/{paper_id}/regrade-sheet", "重新批改"],
            ["GET", "/api/exam-bank/geogebra/animations", "GeoGebra 动画"],
            ["GET", "/api/exam-bank/questions/{question_id}/animations", "题目关联动画"],
            ["POST", "/api/exam-bank/geogebra/sync", "同步 GeoGebra 数据"],
        ],
    )

    add_heading(doc, "2.3 App 客户端（/api/app/*）", 2)
    add_table(
        doc,
        ["方法", "路径", "认证", "用途"],
        [
            ["GET", "/api/app/config", "无", "运行时配置（API/CDN/KaTeX/支付）"],
            ["POST", "/api/app/auth/sms", "无", "发送验证码"],
            ["POST", "/api/app/auth/login", "无", "登录，返回 JWT"],
            ["GET", "/api/app/me", "JWT", "当前用户与会员状态"],
            ["POST", "/api/app/orders", "JWT", "创建支付订单"],
            ["GET", "/api/app/orders/{order_id}", "JWT", "查询订单"],
            ["POST", "/api/app/orders/{order_id}/mock-pay", "JWT", "开发环境模拟支付"],
            ["GET", "/api/app/pay/mock/{order_id}", "token 参数", "模拟支付页"],
            ["POST", "/api/app/pay/mock/{order_id}", "token 参数", "模拟支付确认"],
            ["POST", "/api/app/pay/wechat/notify", "微信签名", "微信支付回调"],
            ["POST", "/api/app/pay/alipay/notify", "支付宝签名", "支付宝回调"],
            ["GET", "/api/app/papers", "JWT + 会员", "试卷列表"],
            ["GET", "/api/app/papers/{paper_id}/preview", "JWT + 会员", "刷题预览 HTML"],
            ["GET", "/api/app/papers/{paper_id}/pdf", "JWT + 会员", "导出 PDF"],
            ["GET", "/api/app/ocr-status", "JWT + 会员", "OCR 状态"],
            ["POST", "/api/app/papers/{paper_id}/grade-sheet", "JWT + 会员", "拍照阅卷"],
            ["GET", "/api/app/knowledge-points", "JWT + 会员", "知识点列表"],
            ["POST", "/api/app/questions-by-knowledge", "JWT + 会员", "按知识点选题"],
            ["POST", "/api/app/practice/preview", "JWT + 会员", "专项训练预览"],
            ["POST", "/api/app/devices", "JWT", "注册推送设备"],
        ],
    )

    add_heading(doc, "2.4 管理端与 App 的对应关系", 2)
    add_table(
        doc,
        ["管理端", "App 等价", "主要差异"],
        [
            ["GET /api/exam-bank/papers", "GET /api/app/papers", "App 需会员；图片可走 CDN"],
            ["GET .../papers/{id}/preview", "GET /api/app/papers/{id}/preview", "App 刷题预览并改写资源地址"],
            ["POST .../questions-by-knowledge", "POST /api/app/questions-by-knowledge", "App 返回字段更精简"],
            ["POST .../practice/preview", "POST /api/app/practice/preview", "App 附带 KaTeX 配置"],
            ["POST .../grade-sheet", "POST /api/app/papers/{id}/grade-sheet", "App 需登录会员"],
            ["GET .../ocr-status", "GET /api/app/ocr-status", "App 需登录会员"],
        ],
    )

    add_heading(doc, "三、相关页面（非 /api）", 1)
    add_table(
        doc,
        ["路径", "归属", "说明"],
        [
            ["/", "管理端", "转换与题库管理"],
            ["/knowledge", "管理端", "知识点维护、上传 PPT"],
            ["/practice", "管理端", "专项训练与 GeoGebra / PPT"],
            ["/view/{job}", "管理端", "转换结果预览"],
            ["/files/{job}/...", "管理端", "临时转换文件"],
            ["/vendor/katex/...", "共用", "KaTeX 静态资源"],
            ["/static/...", "管理端", "可视化解题、课件浏览"],
            ["/health", "运维", "健康检查"],
        ],
    )

    add_heading(doc, "四、数据库核心表", 1)
    add_table(
        doc,
        ["模块", "表"],
        [
            ["题库", "papers、questions、question_options、question_types"],
            ["知识点", "knowledge_points、question_knowledge_points、knowledge_geogebra_animations"],
            ["课件", "knowledge_point_courseware（文件在 Storage exam-assets）"],
            ["资源", "assets"],
            ["答题卡", "answer_sheets、answer_sheet_items"],
            ["App 用户", "app_users、app_memberships、app_devices、app_orders"],
        ],
    )
    add_para(doc, "服务端口：管理端与 API 均为 3010。App 默认连接 http://127.0.0.1:3010（Android 模拟器为 http://10.0.2.2:3010）。")

    doc.save(OUT)
    png.unlink(missing_ok=True)
    print(OUT)


if __name__ == "__main__":
    main()
