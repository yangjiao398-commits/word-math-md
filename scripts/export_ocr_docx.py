"""Write the answer-sheet OCR technical note as a Word file."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
PNG = ROOT / "docs" / "ocr-pipeline.png"
DOC = ROOT / "docs" / "OCR识别技术说明.docx"
FONT = r"C:\Windows\Fonts\msyh.ttc"
FONTB = r"C:\Windows\Fonts\msyhbd.ttc"


def font(size, bold=False):
    return ImageFont.truetype(FONTB if bold else FONT, size)


def set_run(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_p(doc, text, size=11, bold=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    set_run(p.add_run(text), size=size, bold=bold)
    return p


def add_h(doc, text, size=16):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run(text), size=size, bold=True, color=(31, 78, 121))
    return p


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd",
        {
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val": "clear",
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color": "auto",
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill": "1F4E79",
        },
    )
    tcPr.append(shd)
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, head in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        set_run(p.add_run(head), size=10, bold=True, color=(255, 255, 255))
        shade_header(cell)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            set_run(p.add_run(value), size=10)
    doc.add_paragraph()
    return table


def draw():
    img = Image.new("RGB", (1600, 420), "#f7f4ee")
    d = ImageDraw.Draw(img)
    d.text((36, 24), "答卷识别流水线", font=font(32, True), fill="#1a2433")
    steps = [
        ("1. 照片", "jpg / png\n最长边 2800\n对比度 1.25"),
        ("2. 选引擎", "通义千问 VL\nOpenAI 视觉\nRapidOCR\nTesseract"),
        ("3. 读文字", "行文本或 JSON\n按从上到下排序"),
        ("4. 按题号切开", "只认 1. 2. 3.\n忽略（1）（2）"),
        ("5. 抽答案", "选择 / 填空 / 解答\n再和标准答案比对"),
    ]
    x = 36
    for i, (title, body) in enumerate(steps):
        d.rounded_rectangle((x, 100, x + 270, 360), radius=16, fill="#e7f0fa", outline="#1f4e79", width=3)
        d.text((x + 16, 118), title, font=font(24, True), fill="#1f4e79")
        y = 168
        for line in body.split("\n"):
            d.text((x + 16, y), line, font=font(20), fill="#243140")
            y += 36
        if i < len(steps) - 1:
            d.line((x + 270, 230, x + 308, 230), fill="#1f4e79", width=4)
            d.polygon([(x + 318, 230), (x + 302, 222), (x + 302, 238)], fill="#1f4e79")
        x += 312
    PNG.parent.mkdir(parents=True, exist_ok=True)
    img.save(PNG)


def document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(title.add_run("答卷 OCR：识别题目答案的技术说明"), size=22, bold=True, color=(26, 36, 51))
    add_p(
        doc,
        "本文说明 word-math-md「识别并评分」现在怎么读答题卡，RapidOCR 还能调哪些参数，以及哪些 OCR 更适合中文手写数学公式。",
        size=11,
    )

    add_h(doc, "一、当前识别流程")
    add_p(
        doc,
        "入口是管理端 POST /api/exam-bank/papers/{id}/grade-sheet，以及 App 的 POST /api/app/papers/{id}/grade-sheet。服务先取出这套卷的题号和题型，再对每一张照片做 OCR，最后按题号顺序抽出学生答案，和 answer_md 比对打分。",
    )
    doc.add_picture(str(PNG), width=Cm(17.2))
    add_p(doc, "图 1  从照片到评分的五步。", size=10)

    add_p(doc, "引擎按这个顺序自动选择：通义千问 VL（DASHSCOPE_API_KEY）、OpenAI 视觉（OPENAI_API_KEY）、RapidOCR、Tesseract。可以用环境变量 SHEET_OCR_ENGINE 指定其中一个。最近几次评分实际走的是 RapidOCR。", size=11)
    add_p(doc, "照片预处理在 word_math_md/sheet_ocr.py：按 EXIF 转正，长边超过 2800 像素才缩小，对比度提高到 1.25 并锐化，然后存成 JPEG（质量 85）或 PNG。单张不超过 12MB，一次最多 12 张。", size=11)

    add_h(doc, "二、RapidOCR 读出文字之后怎么变成答案")
    add_p(
        doc,
        "RapidOCR 返回的是一行行文字和检测框，不是「第几题选了什么」。程序按框的纵向中心从上到下排成一篇文本，再交给 parse_answers_in_order。视觉模型则先按提示词返回 JSON；若某个选择题答案不是纯字母，就退回同一套按题号抽取的规则。",
    )
    add_table(
        doc,
        ["步骤", "规则", "避开的错误"],
        [
            ["切大题", "只把行首的 1. / 2. / 第3题 当作大题", "解答题里的（1）（2）不再被当成第 1、2 题"],
            ["单选", "读「为（D）」「是（C」「（）C」或单独一行的字母", "不把印刷选项 A. B. C. D. 当成学生答案"],
            ["多选", "同样位置上的连续字母，如（BD）或单独一行 ABD", "不把题干汉字收进来"],
            ["填空", "只取本题最后一个「为 / =」后面的内容，含写在括号里的式子", "不把前面的「圆心为 O」当成答案"],
            ["解答", "（1）（2）之后的手写行拼在一起", "题干和选项行不计入"],
            ["比对", "选择比字母；填空把分数和 π 收成同一写法", "2π/3+2 与标准答案 (2/3)π+2 视为相同"],
            ["未读到", "该题没有抽出答案", "结果显示「未识别」，不再标成「错误 · 未识别」"],
        ],
    )
    add_p(
        doc,
        "括号重叠是文字规则补上的，不是 RapidOCR 自己懂数学。它如果把「（2π/3+2）」读成「面积为」，后面的规则没有字可抽，这题就仍是未识别。",
    )

    add_h(doc, "三、RapidOCR 提升准确度的手段")
    add_p(
        doc,
        "本机安装的是 rapidocr_onnxruntime，模型是 PaddleOCR 的 PP-OCRv3：检测 ch_PP-OCRv3_det，方向分类 ch_ppocr_mobile_v2.0_cls，识别 ch_PP-OCRv3_rec。它擅长印刷中文和较工整的数字、字母，输出普通文字，不输出 LaTeX。",
    )
    add_p(doc, "官方调优顺序是：先放宽检测框，再放宽识别置信度，再放大文本框；这三样仍不够，就换更大的检测/识别模型。本项目已经动过前三样，调用时传入：", size=11)
    add_table(
        doc,
        ["参数", "默认", "本项目现在", "作用"],
        [
            ["box_thresh", "0.5", "0.3", "检测框要有多像文字才保留。调低能找回单独的小字，也会多出噪声"],
            ["text_score", "0.5", "0.3", "识别结果的置信度门槛。调低能留下模糊笔迹，错字也会变多"],
            ["unclip_ratio", "1.6", "2.2", "把检测框向外撑大。写在括号上的笔迹更不容易被裁掉"],
            ["最长边", "—", "2800 像素", "缩小太狠时，括号里的字会低于检测模型的有效高度"],
            ["对比度 / 锐化", "—", "1.25 + 锐化", "让浅色笔迹和纸面拉开一点"],
        ],
    )
    add_p(doc, "配置文件里还有这些旋钮，当前调用没有改：", size=11)
    add_table(
        doc,
        ["参数", "当前配置", "什么时候动"],
        [
            ["thresh", "0.3", "文字和背景的分割阈值。字被背景吃掉时再略调低"],
            ["min_height", "30", "矮于 30 像素的框会走整图识别。括号里的单字母经常更矮，可试 15–20"],
            ["limit_side_len", "736，limit_type=min", "检测前把短边放到 736。整页答卷上的小字仍偏小，可试 960 或 1280，速度会下降"],
            ["use_dilation", "true", "检测区域做膨胀，细笔划更容易连成一块。已开启"],
            ["score_mode", "fast", "改成 slow 时框的得分更稳，检测更慢"],
            ["use_angle_cls", "true", "只区分 0° 和 180°。侧着拍的卷子要先转正，分类器不会扶正 90°"],
            ["max_side_len", "新版 RapidOCR 默认 2000", "本包装的是旧版接口，长边由项目自己限制在 2800"],
        ],
    )
    add_p(doc, "参数仍解决不了时，RapidOCR 文档给出的下一步是换模型，而不是继续把阈值打到很低：", size=11)
    add_table(
        doc,
        ["手段", "说明"],
        [
            ["换 server 检测模型", "先换检测。漏框、粘连、小字找不到时，server 检测比 mobile 稳"],
            ["再换 server 识别模型", "框已经套准、字仍认错时再换识别模型"],
            ["升级到 PP-OCRv4 / v5", "新版 RapidOCR 包可选 mobile 或 server。本仓库用的 onnxruntime 包装仍是 PP-OCRv3"],
            ["PP-OCRv6", "2026 年文档中的新版本，量级有 tiny / small / medium。仍是通用文字模型，不是公式模型"],
            ["图像补边", "文字贴着照片边缘时，四周补白再识别，避免框被切掉"],
            ["不要用整页硬认公式", "PP-OCR 的识别头是字符序列。叠分式、根号、上下标会被摊成一行，调阈值改变不了这件事"],
        ],
    )

    add_h(doc, "四、更适合中文手写数学公式的 OCR")
    add_p(
        doc,
        "选择题字母用行文字 OCR 就够。填空和解答里的分式、根号、π、上下标，需要能输出 LaTeX 的公式识别，或能看完整页的视觉模型。下表里的中文指标来自 PaddleOCR 公布的公式测试集，不是这套答题卡上的实测。",
    )
    add_table(
        doc,
        ["方案", "适合什么", "输出", "和本项目的关系"],
        [
            [
                "PP-FormulaNet_plus-M / L",
                "中文试卷、教材、论文里的公式。Paddle 建议中文场景用这两档。plus-M 的中文 BLEU 约 89.76，plus-L 约 90.64",
                "LaTeX",
                "适合裁出填空空白后再识别，不必替换整页 RapidOCR",
            ],
            [
                "UniMERNet",
                "训练数据含印刷、扫描和手写公式。中文 BLEU 约 43.50，低于 plus 系列，但手写覆盖是明确写进训练说明的",
                "LaTeX",
                "手写公式可做对照；整页中文题干仍要另做检测",
            ],
            [
                "Pix2Text（mfd-1.5 + mfr-1.5）",
                "整页版面：中文用 CnOCR，公式单独检测再识别，最后合成 Markdown。定位是本地的 Mathpix 替代",
                "Markdown + LaTeX",
                "最接近「一页答卷直接变成题干和公式」",
            ],
            [
                "LaTeX-OCR",
                "英文印刷公式。中文 BLEU 约 39.96",
                "LaTeX",
                "不适合中文手写答卷",
            ],
            [
                "通义千问 VL / GPT-4o",
                "整页里区分印刷题干和学生笔迹，尤其是括号里的选项字母",
                "按提示词的 JSON",
                "项目已接入，且排在 RapidOCR 前面。有密钥时会优先使用",
            ],
            [
                "Mathpix",
                "商业接口，印刷和手写公式、化学式、整份 PDF 都覆盖",
                "LaTeX / Markdown",
                "准确度高，按量计费，答卷会离开本机",
            ],
        ],
    )
    add_p(
        doc,
        "Paddle 自己的建议是：中文公式用 PP-FormulaNet_plus-L 或 plus-M；英文且要速度时用较小的 S 模型。UniMERNet 更值得在「确认是手写公式」时做对比，不要只看它在中文印刷公式上的 BLEU。",
    )

    add_h(doc, "五、对本项目的用法")
    add_table(
        doc,
        ["题型", "现在", "若还要提高"],
        [
            ["单选、多选", "RapidOCR 行文本 + 括号内字母规则", "笔迹压住括号时，优先用已接入的视觉模型，或把空白裁出来再认"],
            ["填空公式", "把最后一个空后面的文字收成 2/3pi+2 这种写法再比对", "空白区域改走 PP-FormulaNet_plus-M 或 Pix2Text 的公式模型，直接得到 LaTeX"],
            ["解答过程", "（1）（2）之后的手写行做粗比对，对不上就待复核", "整页用 Pix2Text 或视觉模型，RapidOCR 只负责题号和选项字母"],
        ],
    )
    add_p(
        doc,
        "RapidOCR 继续负责「看见字」。公式结构交给公式模型。两者叠在同一次评分里：题号和选择题仍走现在的切分，只有填空和解答的空白图块改走公式 OCR。",
    )
    add_p(doc, "相关代码：word_math_md/sheet_ocr.py、word_math_md/sheet_grade.py。", size=10)
    doc.save(DOC)


def main():
    draw()
    document()
    print(DOC)


if __name__ == "__main__":
    main()
