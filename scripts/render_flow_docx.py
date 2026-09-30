"""Render the Word-to-App mermaid flowchart into a Word document."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
PNG = ROOT / "docs" / "word-to-app-flow.png"
DOC = ROOT / "docs" / "从Word上传到App刷题-流程图.docx"
FONT = r"C:\Windows\Fonts\msyh.ttc"
FONTB = r"C:\Windows\Fonts\msyhbd.ttc"


def F(size, bold=False):
    return ImageFont.truetype(FONTB if bold else FONT, size)


def text_size(draw, text, fnt):
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def wrap_center(draw, xy, lines, fnt, fill="#102033"):
    x0, y0, x1, y1 = xy
    heights = [text_size(draw, line, fnt)[1] for line in lines]
    gap = 4
    total = sum(heights) + gap * (len(lines) - 1)
    y = y0 + (y1 - y0 - total) / 2
    for line, h in zip(lines, heights):
        w, _ = text_size(draw, line, fnt)
        draw.text(((x0 + x1 - w) / 2, y), line, font=fnt, fill=fill)
        y += h + gap


def round_box(draw, xy, lines, fill, outline="#24527a", fnt=None):
    draw.rounded_rectangle(xy, radius=16, fill=fill, outline=outline, width=3)
    wrap_center(draw, xy, lines, fnt or F(22, True))


def cylinder(draw, xy, lines, fill="#fff4d6"):
    x0, y0, x1, y1 = xy
    eh = 28
    draw.ellipse((x0, y1 - eh * 2, x1, y1), fill=fill, outline="#8a6a12", width=3)
    draw.rectangle((x0, y0 + eh, x1, y1 - eh), fill=fill, outline=fill)
    draw.line((x0, y0 + eh, x0, y1 - eh), fill="#8a6a12", width=3)
    draw.line((x1, y0 + eh, x1, y1 - eh), fill="#8a6a12", width=3)
    draw.ellipse((x0, y0, x1, y0 + eh * 2), fill=fill, outline="#8a6a12", width=3)
    wrap_center(draw, (x0 + 8, y0 + eh, x1 - 8, y1 - 8), lines, F(20, True))


def group(draw, xy, title, fill="#f4f8fb", outline="#7aa0c4"):
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline=outline, width=3)
    draw.text((xy[0] + 22, xy[1] + 14), title, font=F(24, True), fill="#1f4e79")


def arrow(draw, start, end, color="#2f6ea8"):
    x0, y0 = start
    x1, y1 = end
    draw.line((x0, y0, x1, y1), fill=color, width=4)
    import math

    ang = math.atan2(y1 - y0, x1 - x0)
    size = 14
    p1 = (x1 - size * math.cos(ang - 0.4), y1 - size * math.sin(ang - 0.4))
    p2 = (x1 - size * math.cos(ang + 0.4), y1 - size * math.sin(ang + 0.4))
    draw.polygon([(x1, y1), p1, p2], fill=color)


def label(draw, xy, text, fill="#1f4e79"):
    fnt = F(16, True)
    w, h = text_size(draw, text, fnt)
    x, y = xy
    draw.rounded_rectangle((x - 6, y - 2, x + w + 6, y + h + 6), radius=6, fill="#ffffff")
    draw.text((x, y), text, font=fnt, fill=fill)


def main():
    W, H = 2400, 1680
    img = Image.new("RGB", (W, H), "#ffffff")
    d = ImageDraw.Draw(img)
    d.text((48, 28), "从 Word 上传到 App 刷题", font=F(40, True), fill="#102033")
    d.text((48, 84), "管理端写入 Supabase，App 只读 /api/app/*", font=F(22), fill="#4a6072")

    group(d, (40, 140, 620, 430), "管理端 Web（http://127.0.0.1:3010）", "#eef6fc")
    round_box(d, (80, 210, 580, 290), ["上传解析版 .docx"], "#d6e8f8", fnt=F(24, True))
    round_box(d, (80, 320, 580, 400), ["或上传 .md"], "#d6e8f8", fnt=F(24, True))

    group(d, (40, 470, 620, 1600), "转换管线（本地 server）", "#eef8f1")
    nodes = {
        "B1": (80, 540, 580, 680),
        "B2": (80, 740, 580, 860),
        "B3": (80, 920, 580, 1100),
        "B4": (80, 1160, 580, 1280),
        "B5": (80, 1340, 580, 1500),
    }
    texts = {
        "B1": ["OLE 公式 → LaTeX", "ole_to_latex.py"],
        "B2": ["Word 预处理", "preprocess_docx"],
        "B3": ["Node/TS 转换", "gaokao_docx/convert.mts", "mammoth + turndown + KaTeX"],
        "B4": ["解析 Markdown", "parse-md.mts"],
        "B5": ["入库 exam_bank.py", "拆题、知识点、图片"],
    }
    for key, xy in nodes.items():
        round_box(d, xy, texts[key], "#e5f6ea", fnt=F(22, True))

    arrow(d, (330, 290), (330, 540))
    label(d, (360, 430), "POST /api/exam-bank/import-docx")
    for a, b in (("B1", "B2"), ("B2", "B3"), ("B3", "B4"), ("B4", "B5")):
        arrow(d, (330, nodes[a][3]), (330, nodes[b][1]))

    # md upload jumps to B5 along the left gutter
    d.line((70, 360, 55, 360, 55, 1420, 80, 1420), fill="#2f6ea8", width=4)
    arrow(d, (70, 1420), (80, 1420))
    label(d, (70, 1488), "POST /api/exam-bank/import-markdown")

    group(d, (680, 470, 1280, 980), "Supabase", "#fffaf0")
    cylinder(d, (730, 560, 1180, 760), ["PostgreSQL", "papers / questions / options", "knowledge_points / assets"])
    cylinder(d, (730, 800, 1180, 960), ["Storage exam-assets", "试卷图片"])

    arrow(d, (580, 1420), (730, 680))
    label(d, (600, 1000), "写入题库")
    arrow(d, (580, 1460), (730, 880))
    label(d, (600, 1180), "上传图片")

    group(d, (1340, 140, 2340, 1600), "App 刷题（math-agent-app）", "#fff5f3")
    app = {
        "C1": (1400, 210, 1860, 290),
        "C2": (1400, 320, 1860, 420),
        "C3": (1400, 450, 1860, 550),
        "C4": (1400, 600, 1860, 700),
        "C5": (1400, 760, 1860, 880),
        "C6": (1400, 980, 1860, 1080),
        "C7": (1920, 450, 2280, 530),
        "C8": (1920, 570, 2280, 670),
        "C9": (1920, 710, 2280, 830),
        "C10": (1920, 870, 2280, 990),
    }
    app_text = {
        "C1": ["GET /api/app/config"],
        "C2": ["POST /api/app/auth/login", "JWT"],
        "C3": ["GET /api/app/me", "校验会员"],
        "C4": ["GET /api/app/papers", "试卷列表"],
        "C5": ["GET /api/app/papers/{id}/preview", "KaTeX HTML"],
        "C6": ["WebView 渲染题目"],
        "C7": ["专项训练路径"],
        "C8": ["GET /api/app/knowledge-points"],
        "C9": ["POST /api/app/questions-by-knowledge"],
        "C10": ["POST /api/app/practice/preview"],
    }
    for key, xy in app.items():
        fill = "#fde8e4" if key != "C7" else "#f7f1ee"
        round_box(d, xy, app_text[key], fill, fnt=F(18, True))
    for a, b in (("C1", "C2"), ("C2", "C3"), ("C3", "C4"), ("C4", "C5"), ("C5", "C6")):
        arrow(d, ((app[a][0] + app[a][2]) // 2, app[a][3]), ((app[b][0] + app[b][2]) // 2, app[b][1]))
    arrow(d, (1860, 500), (1920, 620))
    for a, b in (("C8", "C9"), ("C9", "C10")):
        arrow(d, (2100, app[a][3]), (2100, app[b][1]))
    arrow(d, (1920, 940), (1860, 1030))

    arrow(d, (1180, 660), (1400, 650))
    label(d, (1188, 600), "supabase-py 读取")
    arrow(d, (1180, 700), (1400, 820))
    label(d, (1190, 760), "supabase-py 读取")
    arrow(d, (1180, 880), (1400, 1030))
    label(d, (1190, 900), "CDN 或 Supabase URL")
    label(d, (1190, 930), "rewrite 到 APP_CDN_BASE")

    img.save(PNG, "PNG")

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.left_margin = Cm(1.2)
    sec.right_margin = Cm(1.2)
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("从 Word 上传到 App 刷题")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(16, 32, 51)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    doc.add_picture(str(PNG), width=Cm(27))
    doc.save(DOC)
    print(DOC)


if __name__ == "__main__":
    main()
