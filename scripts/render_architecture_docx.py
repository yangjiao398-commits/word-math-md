"""Architecture overview diagram and Word document for word-math-md."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
PNG = ROOT / "docs" / "architecture-overview.png"
DOC = ROOT / "docs" / "架构总览.docx"
FONT = r"C:\Windows\Fonts\msyh.ttc"
FONTB = r"C:\Windows\Fonts\msyhbd.ttc"


def F(size, bold=False):
    return ImageFont.truetype(FONTB if bold else FONT, size)


def box(draw, xy, title, lines, fill, outline):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    draw.text((xy[0] + 18, xy[1] + 14), title, font=F(26, True), fill=outline)
    y = xy[1] + 58
    for line in lines:
        draw.text((xy[0] + 18, y), line, font=F(20), fill="#243140")
        y += 32


def arrow(draw, a, b, color="#3d6ea8"):
    draw.line((*a, *b), fill=color, width=4)
    import math

    ang = math.atan2(b[1] - a[1], b[0] - a[0])
    s = 14
    p1 = (b[0] - s * math.cos(ang - 0.4), b[1] - s * math.sin(ang - 0.4))
    p2 = (b[0] - s * math.cos(ang + 0.4), b[1] - s * math.sin(ang + 0.4))
    draw.polygon([b, p1, p2], fill=color)


def set_font(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = RGBColor(*color)


def heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run, 16 if level == 1 else 14, True, (31, 78, 121))


def para(doc, text):
    p = doc.add_paragraph()
    set_font(p.add_run(text), 11)


def shade(row):
    for cell in row.cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(tcPr.makeelement(qn("w:shd"), {qn("w:fill"): "1F4E79", qn("w:val"): "clear"}))


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    shade(t.rows[0])
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            t.rows[r].cells[c].text = value
            for p in t.rows[r].cells[c].paragraphs:
                for run in p.runs:
                    set_font(run, 10)
    doc.add_paragraph()


def draw():
    img = Image.new("RGB", (2200, 1280), "#ffffff")
    d = ImageDraw.Draw(img)
    d.text((40, 24), "word-math-md 架构总览", font=F(40, True), fill="#102033")
    d.text((40, 80), "管理端与 App 都连 3010；只有后台读写 Supabase", font=F(22), fill="#4a6072")

    box(
        d, (40, 150, 560, 620),
        "前端 · 管理端 Web",
        ["FastAPI 内嵌 HTML / CSS / JS", "页面：/  /knowledge  /practice", "KaTeX 0.18 渲染公式", "visual-solve.js 本地 SVG 动画", "GeoGebra Web（人教配套课件）", "Office Online 预览 PPT"],
        "#eef6fc", "#1f4e79",
    )
    box(
        d, (820, 150, 1480, 760),
        "后台 · Python 服务 :3010",
        ["FastAPI + Uvicorn + Pydantic", "Word：python-docx / mammoth / lxml", "公式：OMML、MTEF、mathml2latex", "Node/TS：tsx、turndown、jszip、linkedom", "题库：exam_bank.py、batch_import.py", "App API：JWT、会员、微信/支付宝", "OCR：RapidOCR / 通义千问 VL / OpenAI", "PDF：Chrome 或 Edge 无头打印"],
        "#eef8f1", "#1d6b3a",
    )
    box(
        d, (1580, 150, 2160, 520),
        "数据库 · Supabase",
        ["PostgreSQL", "Storage 桶 exam-assets", "supabase-py + Service Role", "表级 RLS"],
        "#fff8e8", "#8a6a12",
    )
    box(
        d, (40, 700, 560, 1160),
        "App · math-agent-app",
        ["Flutter / Dart ≥ 3.3", "Material 3 + provider", "http 调用 /api/app/*", "WebView + KaTeX 看题", "image_picker 拍照阅卷", "qr_flutter 支付二维码"],
        "#fff5f3", "#8a3b32",
    )
    box(
        d, (820, 840, 1480, 1160),
        "外部服务",
        ["GeoGebra API / geogebra.org", "微信、支付宝回调", "通义千问 VL、OpenAI Vision", "Office Online 预览"],
        "#f4f1fb", "#5b4b8a",
    )

    arrow(d, (560, 380), (820, 380))
    d.text((600, 340), "/api/exam-bank", font=F(16, True), fill="#1f4e79")
    arrow(d, (300, 700), (860, 760))
    d.text((430, 680), "/api/app/*", font=F(16, True), fill="#8a3b32")
    arrow(d, (1480, 340), (1580, 340))
    d.text((1490, 300), "读写", font=F(16, True), fill="#8a6a12")
    arrow(d, (1150, 760), (1150, 840))
    d.text((1160, 788), "调用", font=F(16, True), fill="#5b4b8a")
    img.save(PNG)


def document():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21)
    sec.left_margin = Cm(1.4)
    sec.right_margin = Cm(1.4)
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("word-math-md 架构总览"), 20, True, (31, 78, 121))
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p2.add_run("前端、后台、数据库、App 的技术工具与关联"), 12, False, (74, 96, 114))

    heading(doc, "一、架构图", 1)
    para(doc, "管理端网页和 Flutter App 都访问本机 3010 端口。入库、会员、支付只由 Python 后台完成，客户端不直连数据库。")
    doc.add_picture(str(PNG), width=Cm(26))

    heading(doc, "二、各层技术工具", 1)
    table(doc, ["层级", "技术工具", "职责"], [
        ["前端（管理端）", "FastAPI 返回的 HTML / CSS / 原生 JavaScript", "转换、题库、知识点、专项训练页面，无 React/Vue"],
        ["前端（公式）", "KaTeX 0.18（npm，挂到 /vendor/katex）", "把 LaTeX 渲染成公式"],
        ["前端（可视化）", "visual-solve.js 本地 SVG；GeoGebra Web", "解题动画是本地脚本；人教配套课件才用 GeoGebra"],
        ["前端（课件）", "Office Online embed", "预览已上传的 PPT"],
        ["后台", "Python 3.10+、FastAPI、Uvicorn、Pydantic、python-dotenv", "全部 HTTP API，默认端口 3010"],
        ["后台（Word）", "python-docx、mammoth、lxml、olefile、Pillow、mathml2latex、BeautifulSoup", "清理 Word、导出图片、OMML/OLE 转 LaTeX"],
        ["后台（可选）", "Aspose.Words", "商业路径，更强的 Word 清理"],
        ["后台（高考解析）", "Node.js + TypeScript：tsx、mammoth、turndown、jszip、linkedom、KaTeX", "解析版 Word 转 Markdown 并拆题"],
        ["后台（阅卷/PDF/支付）", "RapidOCR；通义千问 VL / OpenAI Vision；Chrome/Edge 无头；微信与支付宝", "拍照评分、导出 PDF、会员收款"],
        ["数据库", "Supabase：PostgreSQL + Storage 桶 exam-assets", "题目、用户、会员、图片和 PPT"],
        ["数据库访问", "supabase-py，Service Role Key，表级 RLS", "后台绕过 RLS；匿名 key 不能直接读表"],
        ["App", "Flutter（Dart ≥ 3.3），独立仓库 math-agent-app", "iOS / Android / Web 学生端"],
        ["App（界面与网络）", "Material 3、provider、http、shared_preferences", "页面状态、登录 token、请求 3010"],
        ["App（看题与支付）", "webview_flutter + KaTeX；image_picker；qr_flutter；share_plus", "公式、拍照、支付码、分享 PDF"],
    ])

    heading(doc, "三、它们怎么连在一起", 1)
    table(doc, ["从", "到", "怎么关联"], [
        ["管理端网页", "后台 :3010", "同一进程。页面用 fetch 调 /api/exam-bank/* 和转换接口，无 JWT"],
        ["Flutter App", "后台 /api/app/*", "默认 http://127.0.0.1:3010；Android 模拟器为 10.0.2.2:3010。JWT + 会员"],
        ["后台", "Supabase PostgreSQL", "试卷、题目、知识点、课件元数据、用户、订单"],
        ["后台", "Supabase Storage", "试卷图片、知识点 PPT，桶名 exam-assets"],
        ["后台", "Node 转换脚本", "子进程调用 gaokao_docx/convert.mts 与 parse-md.mts"],
        ["管理端 / App", "KaTeX 静态文件", "后台把 node_modules/katex 挂到 /vendor/katex"],
        ["专项训练页", "GeoGebra", "浏览器加载 geogebra.org 脚本；后台用 api.geogebra.org 同步目录"],
        ["后台", "支付与 OCR", "微信、支付宝回调本服务；阅卷可走视觉模型"],
    ])

    heading(doc, "四、一条典型链路", 1)
    para(doc, "管理员在网页上传解析版 Word → FastAPI 调 Python OLE/预处理，再调 Node 转成 Markdown → 题目写入 PostgreSQL，图片写入 Storage。学生用 App 登录拿到 JWT，会员校验通过后请求 /api/app/papers 与预览接口，WebView 用 KaTeX 显示题目。")
    para(doc, "App 不负责导入 Word，也不直接访问 Supabase。")
    doc.save(DOC)


if __name__ == "__main__":
    draw()
    document()
    print(DOC)
